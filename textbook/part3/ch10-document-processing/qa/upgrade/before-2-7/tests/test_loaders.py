import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from document_processing import (
    detect_scanned_pdf,
    load_csv,
    load_docx,
    load_hwpx,
    load_pdf,
    load_pdf_tables_pdfplumber,
    normalize_metadata,
    recognize_hwp_kind,
)

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT.parents[1] / "data" / "rag-corpus"


def write_csv(path: Path) -> None:
    path.write_text(
        '\ufeffname,note\r\nAlice,"first line\nsecond, quoted line"\r\nBob,plain\r\n',
        encoding="utf-8",
    )


def write_hwpx(path: Path, sections: dict[str, str]) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("Contents/content.hpf", "<package/>")
        for name, text in sections.items():
            archive.writestr(
                f"Contents/{name}.xml",
                f'<hp:section xmlns:hp="urn:test"><hp:p><hp:t>{text}</hp:t></hp:p></hp:section>',
            )


def test_normalize_metadata_returns_common_schema_and_scalar_values() -> None:
    result = normalize_metadata(
        {"doc_id": "한글-01", "page": "2", "tags": ["표"]},
        source="안내.PDF",
        file_type=".PDF",
        parser="test",
    )

    assert result == {
        "doc_id": "한글-01",
        "source": "안내.PDF",
        "file_type": "pdf",
        "parser": "test",
        "page": 2,
        "tags": "['표']",
    }


@pytest.mark.parametrize("page", [0, -1])
def test_normalize_metadata_rejects_nonpositive_page(page: int) -> None:
    with pytest.raises(ValueError, match="1 이상"):
        normalize_metadata({"page": page}, source="a.pdf", file_type="pdf", parser="test")


def test_load_pdf_returns_every_text_page_from_actual_corpus() -> None:
    docs = load_pdf(CORPUS / "pdf" / "easy" / "E01.pdf")

    assert len(docs) == 5
    assert [doc.metadata["page"] for doc in docs] == [1, 2, 3, 4, 5]
    assert all(doc.metadata["doc_id"] == "E01" for doc in docs)


def test_load_pdf_tables_extracts_actual_table_cells() -> None:
    docs = load_pdf_tables_pdfplumber(CORPUS / "pdf" / "easy" / "E01.pdf")
    table_pages = [doc for doc in docs if doc.metadata["table_count"] > 0]

    assert [doc.metadata["page"] for doc in table_pages] == [2, 4]
    assert "평일" in table_pages[0].page_content


def test_detect_scanned_pdf_reports_all_image_only_pages() -> None:
    report = detect_scanned_pdf(CORPUS / "pdf" / "scanned" / "S01-E01-scan.pdf")

    assert report.total_pages == 5
    assert report.pages_without_text == (1, 2, 3, 4, 5)
    assert report.likely_scanned


def test_load_docx_preserves_interleaved_paragraph_and_table_order(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    source = DocxDocument()
    source.add_paragraph("표 앞")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "지역"
    table.cell(0, 1).text = "제주"
    source.add_paragraph("표 뒤")
    source.save(path)

    content = load_docx(path)[0].page_content

    assert content.index("표 앞") < content.index("지역 | 제주") < content.index("표 뒤")


def test_load_csv_removes_utf8_bom_from_header(tmp_path: Path) -> None:
    path = tmp_path / "sample.csv"
    write_csv(path)

    first = load_csv(path)[0]

    assert first.page_content.startswith("name: Alice")


def test_load_csv_preserves_quoted_multiline_and_comma_field(tmp_path: Path) -> None:
    path = tmp_path / "sample.csv"
    write_csv(path)

    first = load_csv(path)[0]

    assert "note: first line\nsecond, quoted line" in first.page_content


def test_load_csv_records_data_row_numbers(tmp_path: Path) -> None:
    path = tmp_path / "sample.csv"
    write_csv(path)

    docs = load_csv(path)

    assert [doc.metadata["row"] for doc in docs] == [2, 3]


def test_load_hwpx_returns_section_labels_and_text(tmp_path: Path) -> None:
    path = tmp_path / "sample.hwpx"
    write_hwpx(path, {"section0": "첫 절", "section1": "둘째 절"})

    docs = load_hwpx(path)

    assert [doc.page_content for doc in docs] == ["첫 절", "둘째 절"]
    assert [doc.metadata["section"] for doc in docs] == [1, 2]
    assert all(doc.metadata["file_type"] == "hwpx" for doc in docs)


def test_load_hwpx_uses_natural_section_number_order(tmp_path: Path) -> None:
    path = tmp_path / "out-of-order.hwpx"
    write_hwpx(path, {"section10": "열한 번째", "section2": "세 번째"})

    docs = load_hwpx(path)

    assert [doc.page_content for doc in docs] == ["세 번째", "열한 번째"]


def test_load_hwpx_rejects_legacy_ole_candidate(tmp_path: Path) -> None:
    path = tmp_path / "legacy.hwp"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"not-enough-to-identify-hwp")

    assert recognize_hwp_kind(path) == "ole-candidate"
    with pytest.raises(ValueError, match="전용 파서"):
        load_hwpx(path)
