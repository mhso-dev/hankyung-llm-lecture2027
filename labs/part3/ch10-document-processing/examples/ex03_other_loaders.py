import csv

from _paths import FIXTURES
from docx import Document as DocxDocument

from document_processing import load_csv, load_docx, load_html


def make_fixtures() -> tuple:
    FIXTURES.mkdir(exist_ok=True)
    docx_path = FIXTURES / "sample.docx"
    document = DocxDocument()
    document.add_heading("배송 안내", level=1)
    document.add_paragraph("제주 배송에는 이틀이 더 필요합니다.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "지역", "추가 일수"
    table.cell(1, 0).text, table.cell(1, 1).text = "제주", "2일"
    document.save(docx_path)

    html_path = FIXTURES / "sample.html"
    html_path.write_text(
        "<html><style>숨김</style><h1>반품 안내</h1><p>수령 후 7일 이내 신청합니다.</p></html>",
        encoding="utf-8",
    )
    csv_path = FIXTURES / "sample.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["상품", "보관"])
        writer.writeheader()
        writer.writerows([{"상품": "쿠키", "보관": "실온"}, {"상품": "케이크", "보관": "냉장"}])
    return docx_path, html_path, csv_path


def main() -> None:
    docx_path, html_path, csv_path = make_fixtures()
    docx_docs, html_docs, csv_docs = load_docx(docx_path), load_html(html_path), load_csv(csv_path)
    print(f"docx_docs={len(docx_docs)} table_preserved={'지역 | 추가 일수' in docx_docs[0].page_content}")
    print(f"html_docs={len(html_docs)} style_removed={'숨김' not in html_docs[0].page_content}")
    print(f"csv_docs={len(csv_docs)} rows={[doc.metadata['row'] for doc in csv_docs]}")
    print(f"unicode_ok={'제주' in docx_docs[0].page_content and '반품' in html_docs[0].page_content}")


if __name__ == "__main__":
    main()

