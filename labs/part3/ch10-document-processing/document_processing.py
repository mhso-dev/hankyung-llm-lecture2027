"""Ch 10 문서 처리 실습의 로더, 청커, 임베딩 캐시.

리트리버 래퍼 없이 각 단계의 입력과 출력을 확인할 수 있는 평범한 함수로 구성한다.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from html import unescape
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import pdfplumber
import tiktoken
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

Scalar = str | int | float | bool | None
Embedder = Callable[[list[str]], list[list[float]]]


def normalize_metadata(
    metadata: dict[str, Any], *, source: str | Path, file_type: str, parser: str
) -> dict[str, Scalar]:
    """서로 다른 로더의 메타데이터를 JSON에 기록 가능한 공통 스키마로 바꾼다."""
    source_path = Path(source)
    normalized: dict[str, Scalar] = {
        "doc_id": str(metadata.get("doc_id") or source_path.stem),
        "source": source_path.name,
        "file_type": file_type.lower().lstrip("."),
        "parser": parser,
    }
    if "page" in metadata and metadata["page"] is not None:
        page = int(metadata["page"])
        if page < 1:
            raise ValueError("page는 1 이상의 정수여야 합니다")
        normalized["page"] = page
    for key, value in metadata.items():
        if key in normalized or key == "page":
            continue
        normalized[key] = value if isinstance(value, (str, int, float, bool, type(None))) else str(value)
    return normalized


def load_pdf(path: Path, *, mode: str = "plain") -> list[Document]:
    """pypdf의 일반 추출과 레이아웃 보존 추출을 같은 스키마로 반환한다."""
    if mode not in {"plain", "layout"}:
        raise ValueError("mode는 plain 또는 layout이어야 합니다")
    reader = PdfReader(str(path))
    docs: list[Document] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text(extraction_mode="layout") if mode == "layout" else page.extract_text()
        if not (text or "").strip():
            continue
        metadata = normalize_metadata(
            {"page": page_no}, source=path, file_type="pdf", parser=f"pypdf-{mode}"
        )
        docs.append(Document(page_content=text or "", metadata=metadata))
    return docs


def load_pdf_tables_pdfplumber(path: Path) -> list[Document]:
    """pdfplumber로 페이지 텍스트와 감지된 표의 셀을 함께 추출한다."""
    docs = []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(layout=True) or ""
            tables = page.extract_tables()
            rendered_tables = [
                "\n".join(" | ".join(cell or "" for cell in row) for row in table)
                for table in tables
            ]
            content = "\n\n".join(part for part in [text, *rendered_tables] if part.strip())
            if content:
                metadata = normalize_metadata(
                    {"page": page_no, "table_count": len(tables)},
                    source=path,
                    file_type="pdf",
                    parser="pdfplumber",
                )
                docs.append(Document(page_content=content, metadata=metadata))
    return docs


@dataclass(frozen=True)
class ScanReport:
    total_pages: int
    pages_without_text: tuple[int, ...]

    @property
    def likely_scanned(self) -> bool:
        return self.total_pages > 0 and len(self.pages_without_text) / self.total_pages >= 0.8


def detect_scanned_pdf(path: Path, *, min_text_chars: int = 20) -> ScanReport:
    """OCR을 가장하지 않고 텍스트 레이어가 없는 쪽을 탐지한다."""
    reader = PdfReader(str(path))
    missing = tuple(
        page_no
        for page_no, page in enumerate(reader.pages, start=1)
        if len(re.sub(r"\s+", "", page.extract_text() or "")) < min_text_chars
    )
    return ScanReport(total_pages=len(reader.pages), pages_without_text=missing)


def load_docx(path: Path) -> list[Document]:
    """DOCX의 문단과 표를 문서 순서대로 읽는다."""
    document = DocxDocument(path)
    blocks = []
    table_no = 0
    for block in document.iter_inner_content():
        if hasattr(block, "rows"):
            table_no += 1
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
            blocks.append(f"[표 {table_no}]\n" + "\n".join(rows))
        elif block.text.strip():
            blocks.append(block.text.strip())
    metadata = normalize_metadata({}, source=path, file_type="docx", parser="python-docx")
    return [Document(page_content="\n\n".join(blocks), metadata=metadata)]


def load_html(path: Path) -> list[Document]:
    """HTML의 script/style을 제외하고 제목과 본문 텍스트를 읽는다."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = "\n".join(s.strip() for s in soup.get_text("\n").splitlines() if s.strip())
    metadata = normalize_metadata({}, source=path, file_type="html", parser="beautifulsoup")
    return [Document(page_content=unescape(text), metadata=metadata)]


def load_csv(path: Path) -> list[Document]:
    """CSV 한 행을 Document 하나로 만들고 행 번호를 보존한다."""
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    docs = []
    for row_no, row in enumerate(rows, start=2):
        text = "\n".join(f"{key}: {value}" for key, value in row.items())
        metadata = normalize_metadata(
            {"row": row_no}, source=path, file_type="csv", parser="csv.DictReader"
        )
        docs.append(Document(page_content=text, metadata=metadata))
    return docs


def recognize_hwp_kind(path: Path) -> str:
    """파일 시그니처로 HWP 이진 형식과 HWPX ZIP 형식을 구분한다."""
    signature = path.read_bytes()[:8]
    if signature == bytes.fromhex("D0CF11E0A1B11AE1"):
        return "ole-candidate"
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
        if "Contents/content.hpf" in names or any(
            name.startswith("Contents/section") and name.endswith(".xml") for name in names
        ):
            return "hwpx"
    return "unknown"


def load_hwpx(path: Path) -> list[Document]:
    """HWPX의 section XML에서 텍스트를 읽는다. 레거시 HWP는 명시적으로 거부한다."""
    kind = recognize_hwp_kind(path)
    if kind == "ole-candidate":
        raise ValueError("OLE 컨테이너 후보는 HWP 여부와 본문을 확인할 전용 파서가 필요합니다")
    if kind != "hwpx":
        raise ValueError("유효한 HWPX 패키지를 확인하지 못했습니다")
    docs: list[Document] = []
    with zipfile.ZipFile(path) as archive:
        sections = sorted(
            (
                name
                for name in archive.namelist()
                if name.startswith("Contents/section") and name.endswith(".xml")
            ),
            key=lambda name: int(re.search(r"section(\d+)\.xml$", name).group(1)),  # type: ignore[union-attr]
        )
        for section_no, name in enumerate(sections, start=1):
            root = ElementTree.fromstring(archive.read(name))
            texts = [node.text or "" for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "t"]
            content = "".join(texts).strip()
            if content:
                metadata = normalize_metadata(
                    {"section": section_no}, source=path, file_type="hwpx", parser="zip-xml"
                )
                docs.append(Document(page_content=content, metadata=metadata))
    return docs


def _with_chunk_metadata(doc: Document, text: str, chunk_no: int, method: str) -> Document:
    metadata = dict(doc.metadata)
    location = next(
        (f"{key}{metadata[key]}" for key in ("page", "section", "row") if key in metadata),
        "unit1",
    )
    metadata.update(
        {
            "chunk_id": f"{metadata.get('doc_id', 'doc')}:{location}:{method}:{chunk_no}",
            "chunk_method": method,
        }
    )
    return Document(page_content=text, metadata=metadata)


def split_char(docs: Sequence[Document], *, chunk_size: int, overlap: int) -> list[Document]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("chunk_size > overlap >= 0이어야 합니다")
    result = []
    for doc in docs:
        start = 0
        for chunk_no in range(math.ceil(max(1, len(doc.page_content)) / (chunk_size - overlap))):
            text = doc.page_content[start : start + chunk_size]
            if not text:
                break
            result.append(_with_chunk_metadata(doc, text, chunk_no, "char"))
            start += chunk_size - overlap
    return result


def split_recursive(docs: Sequence[Document], *, chunk_size: int, overlap: int) -> list[Document]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("chunk_size > overlap >= 0이어야 합니다")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", "다. ", ". ", " ", ""],
    )
    result = []
    for doc in docs:
        for chunk_no, text in enumerate(splitter.split_text(doc.page_content)):
            result.append(_with_chunk_metadata(doc, text, chunk_no, "recursive"))
    return result


def split_token(
    docs: Sequence[Document], *, chunk_tokens: int, overlap_tokens: int, encoding: str = "cl100k_base"
) -> list[Document]:
    if chunk_tokens <= 0 or overlap_tokens < 0 or overlap_tokens >= chunk_tokens:
        raise ValueError("chunk_tokens > overlap_tokens >= 0이어야 합니다")
    codec = tiktoken.get_encoding(encoding)
    if any(
        len(codec.encode(character)) > chunk_tokens
        for doc in docs
        for character in set(doc.page_content)
    ):
        raise ValueError("한 글자의 토큰 수가 chunk_tokens보다 커서 유니코드를 보존할 수 없습니다")
    splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
        encoding_name=encoding,
        chunk_size=chunk_tokens,
        chunk_overlap=overlap_tokens,
        separators=["\n\n", "\n", "다. ", ". ", " ", ""],
    )
    result = []
    for doc in docs:
        for chunk_no, text in enumerate(splitter.split_text(doc.page_content)):
            result.append(_with_chunk_metadata(doc, text, chunk_no, "token"))
    return result


def split_structural(docs: Sequence[Document], *, max_chars: int = 1000) -> list[Document]:
    """마크다운 제목과 빈 줄 경계를 우선 보존한 뒤 큰 블록만 재귀 분할한다."""
    if max_chars <= 0:
        raise ValueError("max_chars는 1 이상이어야 합니다")
    result = []
    for doc in docs:
        blocks = re.split(r"(?=^#{1,6}\s)|\n\s*\n", doc.page_content, flags=re.MULTILINE)
        pieces: list[str] = []
        current = ""
        for raw in blocks:
            block = raw.strip()
            if not block:
                continue
            candidate = f"{current}\n\n{block}".strip()
            if current and len(candidate) > max_chars:
                pieces.extend(
                    part.page_content
                    for part in split_recursive(
                        [Document(page_content=current, metadata=doc.metadata)],
                        chunk_size=max_chars,
                        overlap=0,
                    )
                )
                current = block
            else:
                current = candidate
        if current:
            pieces.extend(
                part.page_content
                for part in split_recursive(
                    [Document(page_content=current, metadata=doc.metadata)],
                    chunk_size=max_chars,
                    overlap=0,
                )
            )
        result.extend(_with_chunk_metadata(doc, text, i, "structural") for i, text in enumerate(pieces))
    return result


def cosine_similarity(left: Sequence[float], right: Sequence[float]) -> float:
    dot = sum(a * b for a, b in zip(left, right, strict=True))
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    return dot / (left_norm * right_norm) if left_norm and right_norm else 0.0


def educational_feature_vectors(texts: list[str], *, dimensions: int = 128) -> list[list[float]]:
    """오프라인 실행용 문자 n-gram 특징 벡터. 실제 임베딩 품질 비교에 사용하지 않는다."""
    vectors = []
    for text in texts:
        vector = [0.0] * dimensions
        compact = re.sub(r"\s+", " ", text.lower())
        for n in (1, 2, 3):
            for i in range(max(0, len(compact) - n + 1)):
                token = compact[i : i + n].encode("utf-8")
                index = int.from_bytes(hashlib.sha256(token).digest()[:4]) % dimensions
                vector[index] += 1.0
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        vectors.append([value / norm for value in vector])
    return vectors


def split_semantic(
    docs: Sequence[Document],
    *,
    embedder: Embedder,
    threshold: float = 0.55,
    max_chars: int = 1200,
) -> list[Document]:
    """문장 사이 의미 유사도가 임계값보다 낮은 곳을 경계로 삼는다."""
    if not math.isfinite(threshold) or not -1 <= threshold <= 1:
        raise ValueError("threshold는 -1 이상 1 이하의 유한한 값이어야 합니다")
    if max_chars <= 0:
        raise ValueError("max_chars는 1 이상이어야 합니다")
    result = []
    for doc in docs:
        sentences = [part.strip() for part in re.split(r"(?<=[.!?다요])\s+|\n+", doc.page_content) if part.strip()]
        if not sentences:
            continue
        vectors = embedder(sentences)
        if len(vectors) != len(sentences):
            raise ValueError("embedder가 문장 수와 다른 수의 벡터를 반환했습니다")
        groups = [[sentences[0]]]
        for index in range(1, len(sentences)):
            if cosine_similarity(vectors[index - 1], vectors[index]) < threshold:
                groups.append([])
            groups[-1].append(sentences[index])
        pieces = []
        for group in groups:
            text = " ".join(group)
            pieces.extend(
                part.page_content
                for part in split_recursive(
                    [Document(page_content=text, metadata=doc.metadata)],
                    chunk_size=max_chars,
                    overlap=0,
                )
            )
        result.extend(_with_chunk_metadata(doc, text, i, "semantic") for i, text in enumerate(pieces))
    return result


@dataclass(frozen=True)
class CacheStats:
    hits: int
    misses: int


class EmbeddingCache:
    """모델·제공자·버전·본문을 모두 키에 포함하는 JSON 파일 캐시."""

    def __init__(self, directory: Path) -> None:
        self.directory = directory

    @staticmethod
    def key(
        text: str,
        *,
        provider: str,
        model: str,
        version: str,
        dimensions: int | None = None,
        input_prefix: str = "",
    ) -> str:
        payload = json.dumps(
            {
                "provider": provider,
                "model": model,
                "version": version,
                "dimensions": dimensions,
                "input_prefix": input_prefix,
                "text": text,
            },
            ensure_ascii=False,
            sort_keys=True,
        )
        return hashlib.sha256(payload.encode()).hexdigest()

    def embed(
        self,
        texts: list[str],
        *,
        provider: str,
        model: str,
        version: str,
        embedder: Embedder,
        dimensions: int | None = None,
        input_prefix: str = "",
    ) -> tuple[list[list[float]], CacheStats]:
        self.directory.mkdir(parents=True, exist_ok=True)
        vectors: list[list[float] | None] = [None] * len(texts)
        missing_indexes = []
        for index, text in enumerate(texts):
            key = self.key(
                text,
                provider=provider,
                model=model,
                version=version,
                dimensions=dimensions,
                input_prefix=input_prefix,
            )
            path = self.directory / f"{key}.json"
            if path.exists():
                vectors[index] = json.loads(path.read_text(encoding="utf-8"))["vector"]
            else:
                missing_indexes.append(index)
        if missing_indexes:
            new_vectors = embedder([input_prefix + texts[index] for index in missing_indexes])
            if len(new_vectors) != len(missing_indexes):
                raise ValueError("embedder가 입력 수와 다른 수의 벡터를 반환했습니다")
            for index, vector in zip(missing_indexes, new_vectors, strict=True):
                vectors[index] = vector
                path = self.directory / (
                    self.key(
                        texts[index],
                        provider=provider,
                        model=model,
                        version=version,
                        dimensions=dimensions,
                        input_prefix=input_prefix,
                    )
                    + ".json"
                )
                path.write_text(json.dumps({"vector": vector}), encoding="utf-8")
        return [vector for vector in vectors if vector is not None], CacheStats(
            hits=len(texts) - len(missing_indexes), misses=len(missing_indexes)
        )


def embed_with_cache(
    texts: list[str],
    *,
    cache_dir: Path,
    provider: str,
    model: str,
    version: str,
    dimensions: int | None,
    input_prefix: str,
    embedder: Embedder,
) -> tuple[list[list[float]], CacheStats]:
    """예제에서 캐시 키의 모든 무효화 조건을 드러내는 함수형 진입점."""
    return EmbeddingCache(cache_dir).embed(
        texts,
        provider=provider,
        model=model,
        version=version,
        dimensions=dimensions,
        input_prefix=input_prefix,
        embedder=embedder,
    )


def estimate_embedding_cost(texts: Iterable[str], *, usd_per_million_tokens: float) -> dict[str, float]:
    codec = tiktoken.get_encoding("cl100k_base")
    tokens = sum(len(codec.encode(text)) for text in texts)
    return {"tokens": float(tokens), "estimated_usd": tokens * usd_per_million_tokens / 1_000_000}


def preflight_corpus(paths: Sequence[Path], *, required: int = 100) -> dict[str, Any]:
    content_hashes = {hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
    return {
        "required": required,
        "available": len(content_hashes),
        "ready": len(content_hashes) >= required,
        "duplicate_files": len(paths) - len(content_hashes),
    }


def top_k(
    query: str, chunks: Sequence[Document], vectors: Sequence[Sequence[float]], embedder: Embedder, k: int
) -> list[dict[str, Any]]:
    query_vector = embedder([query])[0]
    ranked = sorted(
        zip(chunks, vectors, strict=True),
        key=lambda item: cosine_similarity(query_vector, item[1]),
        reverse=True,
    )[:k]
    return [
        {
            "doc_id": doc.metadata["doc_id"],
            "page": doc.metadata.get("page"),
            "chunk_id": doc.metadata["chunk_id"],
            "score": round(cosine_similarity(query_vector, vector), 4),
        }
        for doc, vector in ranked
    ]
